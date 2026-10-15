"""Send the Chinese pitch email to the Chinese media list (xlsx) plus test addresses. Simplified for mainland, Traditional for HK/TW."""
import sys, os, json, time, html
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gmail_sender as G

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "chinese_media_covered_alef_with_contacts.xlsx")
STATE = os.path.join(HERE, "cn_send_state.json")
IMG = os.path.join(HERE, "alef_over_tesla_800.jpg"); CID = "alef_over_tesla_800.jpg"
PT = ZoneInfo("America/Los_Angeles")

SUBJ = {"S": "Alef 飞行汽车 vs 特斯拉 Roadster 新闻稿（请于特斯拉发布会后刊发）",
        "T": "Alef 飛行汽車 vs 特斯拉 Roadster 新聞稿（請於特斯拉發布會後刊發）"}

BODY = {"S": """{greet}

感谢贵刊此前对 Alef Aeronautics（Alef 飞行汽车）的报道。我们刚刚准备了一篇重要新闻稿，希望第一时间与您分享。该稿件设有禁发期，请于特斯拉 Roadster 发布会结束后再刊发。

据悉，马斯克将于美国时间 10 月 15 日（北京时间 10 月 16 日上午）发布新款、可飞行的特斯拉 Roadster。如果属实，它将成为继 Alef 之后历史上第二款飞行汽车。

Alef 已提前准备好新闻稿，内容包括：
1. Alef 飞行汽车飞越特斯拉 Roadster 的视频
2. Alef（第一款飞行汽车）与特斯拉 Roadster（第二款飞行汽车）的对比：可在公路行驶 + 垂直起飞

稿件禁发期至特斯拉 Roadster 发布会结束为止。如您同意遵守禁发期并希望获取新闻稿和视频，请回复告知，我们将立即发送中文版新闻稿、视频及高清图片。

谢谢！
Alef 媒体团队 Anna
press@alef.aero

附：Alef 首席执行官吉姆·杜霍夫尼（Jim Dukhovny）可在 10 月 15 日当天及之后随时接受采访，也可提供书面问答。
""",
"T": """{greet}

感謝貴刊此前對 Alef Aeronautics（Alef 飛行汽車）的報導。我們剛剛準備了一篇重要新聞稿，希望第一時間與您分享。該稿件設有禁發期，請於特斯拉 Roadster 發布會結束後再刊發。

據悉，馬斯克將於美國時間 10 月 15 日（台北／香港時間 10 月 16 日上午）發布新款、可飛行的特斯拉 Roadster。如果屬實，它將成為繼 Alef 之後歷史上第二款飛行汽車。

Alef 已提前準備好新聞稿，內容包括：
1. Alef 飛行汽車飛越特斯拉 Roadster 的影片
2. Alef（第一款飛行汽車）與特斯拉 Roadster（第二款飛行汽車）的對比：可在公路行駛 + 垂直起飛

稿件禁發期至特斯拉 Roadster 發布會結束為止。如您同意遵守禁發期並希望取得新聞稿和影片，請回覆告知，我們將立即發送中文版新聞稿、影片及高解析度圖片。

謝謝！
Alef 媒體團隊 Anna
press@alef.aero

附：Alef 執行長吉姆·杜霍夫尼（Jim Dukhovny）可在 10 月 15 日當天及之後隨時接受採訪，也可提供書面問答。
"""}
CAPTION = {"S": "真实照片（非 AI 生成）：Alef 飞行汽车飞越特斯拉 Roadster 与 Cybertruck",
           "T": "真實照片（非 AI 生成）：Alef 飛行汽車飛越特斯拉 Roadster 與 Cybertruck"}
F = "font-family:'Microsoft YaHei','PingFang SC','Noto Sans CJK SC',Arial,sans-serif;font-size:14px;line-height:1.6;color:#000"

def to_html(text, variant):
    paras = text.strip().split("\n\n")
    out = []
    for p in paras:
        lines = [html.escape(l) for l in p.split("\n")]
        body = "<br>".join(lines)
        if p.startswith("附："): body = f"<i>{body}</i>"
        if p.startswith("Alef 已") or p.startswith("Alef 已提前"):
            body = body.replace("1. ", "1. <b>", 1).replace("的视频", "的视频</b>", 1).replace("的影片", "的影片</b>", 1)
        if "马斯克将于" in p or "馬斯克將於" in p:
            body = "<b>" + body + "</b>"
        out.append(f'<p style="margin:0 0 1em 0;{F}">{body}</p>')
    out.append(f'<p style="margin:1.5em 0 0 0;{F}"><img src="cid:{CID}" width="600" style="display:block;max-width:100%;height:auto;border:0" alt="Alef flying over Tesla Roadster"><br><span style="font-size:12px;color:#444">{html.escape(CAPTION[variant])}</span></p>')
    return "<div>" + "\n".join(out) + "</div>"

def greeting(outlet_cn, email_type, region, variant):
    if email_type == "reporter" and outlet_cn == "中国日报": return "朱文倩老师，您好："
    outlet_cn = outlet_cn.split(" /")[0].split(" (")[0].strip()
    if variant == "T": return f"{outlet_cn} 編輯部的老師，您好："
    return f"{outlet_cn} 编辑部的老师，您好："

def recipients():
    ws = load_workbook(XLSX).active
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        outlet_cn, region, email, etype = r[1], r[3], r[8], r[9]
        variant = "T" if region in ("Hong Kong", "Taiwan") else "S"
        rows.append({"email": email.strip(), "outlet": outlet_cn, "variant": variant, "greet": greeting(outlet_cn, etype, region, variant)})
    tests = [{"email": "jwizard@gmail.com", "outlet": "TEST", "variant": "S", "greet": "Jim 老师，您好："},
             {"email": "wizard_files@yahoo.com", "outlet": "TEST", "variant": "T", "greet": "Jim 老師，您好："}]
    return tests + rows

def main():
    dry = "--dry" in sys.argv
    st = json.load(open(STATE)) if os.path.exists(STATE) else {}
    for i, r in enumerate(recipients()):
        if st.get(r["email"], {}).get("status") == "sent": continue
        text = BODY[r["variant"]].format(greet=r["greet"]) + "\n[" + CAPTION[r["variant"]] + "]\n"
        msg = G.build(r["email"], SUBJ[r["variant"]], text, to_html(BODY[r["variant"]].format(greet=r["greet"]), r["variant"]), image_path=IMG, cid=CID)
        if dry:
            print("DRY", r["email"], r["outlet"], r["variant"], r["greet"]); continue
        if i: time.sleep(15)
        try:
            res = G.send(msg); st[r["email"]] = {"status": "sent", "id": res.get("id"), "outlet": r["outlet"], "at": datetime.now(PT).isoformat()}
        except Exception as e:
            st[r["email"]] = {"status": "error", "error": str(e)[:300], "outlet": r["outlet"], "at": datetime.now(PT).isoformat()}
        json.dump(st, open(STATE, "w"), indent=1, ensure_ascii=False)
        print(datetime.now(PT).strftime("%H:%M:%S"), r["email"], r["outlet"], st[r["email"]]["status"], flush=True)
    print("sent:", sum(1 for v in st.values() if v["status"] == "sent"), "errors:", sum(1 for v in st.values() if v["status"] == "error"))

if __name__ == "__main__": main()
