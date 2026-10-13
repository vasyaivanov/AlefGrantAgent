"""Press blast runner. Never stops on bounces or single-send errors.
Usage: python3 run_campaign.py send [--now] [--limit N] [--dry]   # send per schedule in send_list_final.json
       python3 run_campaign.py report                              # build report xlsx + email it
       python3 run_campaign.py bounces                             # scan for bounces since start, email final report
"""
import sys, os, json, time, re, traceback
from datetime import datetime
from zoneinfo import ZoneInfo
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gmail_sender as G, email_template as T
from email.utils import getaddresses

HERE = os.path.dirname(os.path.abspath(__file__))
LIST = os.path.join(HERE, "send_list_final.json")
STATE = os.path.join(HERE, "run_state.json")
PT = ZoneInfo("America/Los_Angeles")
ME = "press@alef.aero"
NEW_SUBJ = T.SUBJECT

def load_state():
    return json.load(open(STATE)) if os.path.exists(STATE) else {"started": None, "rows": {}}
def save_state(s): json.dump(s, open(STATE, "w"), indent=1)

def recipients(r):
    """Reply All: original sender + others on their last message, minus us, minus anyone already emailed elsewhere (precomputed as r['cc'])."""
    if r["mode"] != "reply": return [r["email"]], []
    return [r["email"]], list(r.get("cc", []))

def subject_for(r):
    if r["mode"] != "reply": return NEW_SUBJ
    s = r.get("subject") or NEW_SUBJ
    return s if s.lower().startswith("re:") else "Re: " + s

def send_one(r, dry=False):
    to, cc = recipients(r)
    html, text = T.build_html(r["first_name"]), T.build_text(r["first_name"])
    m = G.build(", ".join(to), subject_for(r), text, html, image_path=os.path.join(HERE, T.IMAGE_FILE), cid=T.IMAGE_CID,
                in_reply_to=r.get("reply_message_id_hdr") or None, references=r.get("references") or None,
                cc=", ".join(cc) if cc else None)
    if dry: return {"dry": True, "to": to, "cc": cc, "subject": subject_for(r)}
    last = None
    for attempt in range(4):
        try:
            return G.send(m, thread_id=r.get("thread_id") or None)
        except Exception as e:
            last = e; msg = str(e)
            if "429" in msg or "500" in msg or "503" in msg: time.sleep(20 * (attempt + 1)); continue
            raise
    raise last

def cmd_send(args):
    now_mode = "--now" in args; dry = "--dry" in args
    limit = int(args[args.index("--limit") + 1]) if "--limit" in args else None
    rows = json.load(open(LIST)); st = load_state()
    if not st["started"]: st["started"] = datetime.now(PT).isoformat(); save_state(st)
    done = 0
    for r in rows:
        key = r["email"]
        if st["rows"].get(key, {}).get("status") == "sent": continue
        if limit is not None and done >= limit: break
        if os.path.exists(os.path.join(HERE, "STOP")): print("STOP file found, halting before", key, flush=True); break
        if not now_mode:
            t = datetime.strptime(r["send_time_pt"], "%Y-%m-%d %I:%M:%S %p").replace(tzinfo=PT)
            wait = (t - datetime.now(PT)).total_seconds()
            while wait > 0:
                if os.path.exists(os.path.join(HERE, "STOP")): print("STOP file found while waiting", flush=True); return
                time.sleep(min(wait, 30)); wait = (t - datetime.now(PT)).total_seconds()
        else:
            time.sleep(15 if done else 0)
        try:
            res = send_one(r, dry=dry)
            st["rows"][key] = {"status": "dry" if dry else "sent", "at": datetime.now(PT).isoformat(), "msg_id": res.get("id"), "mode": r["mode"], "batch": r["batch"]}
        except Exception as e:
            st["rows"][key] = {"status": "error", "at": datetime.now(PT).isoformat(), "error": str(e)[:300], "mode": r["mode"], "batch": r["batch"]}
            print("ERROR", key, e, flush=True)
        save_state(st); done += 1
        print(datetime.now(PT).strftime("%H:%M:%S"), r["seq"], key, st["rows"][key]["status"], flush=True)
    print("done; sent:", sum(1 for v in st["rows"].values() if v["status"] == "sent"), "errors:", sum(1 for v in st["rows"].values() if v["status"] == "error"))

def build_report(final=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    rows = json.load(open(LIST)); st = load_state()
    wb = Workbook(); ws = wb.active; ws.title = "Results"
    ws.append(["#", "Batch", "Planned time (PT)", "Sent at (PT)", "Greeting", "Name", "Outlet", "Email", "Mode", "Status", "Detail"])
    counts = {"sent": 0, "error": 0, "bounced": 0, "not sent": 0}
    for r in rows:
        s = st["rows"].get(r["email"], {}); status = s.get("status", "not sent")
        if status == "dry": status = "not sent"
        counts[status] = counts.get(status, 0) + 1
        ws.append([r["seq"], r["batch"], r["send_time_pt"], s.get("at", "")[:19].replace("T", " "), f"Hi {r['first_name']}." if r["first_name"] else "Hi.",
                   r["name_sheet"], r["outlet"], r["email"], r["mode"], status, s.get("error", "") or s.get("bounce", "")])
    b = Font(name="Arial", size=10, bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]: c.font = b; c.fill = fill
    for col, w in zip("ABCDEFGHIJK", [5, 6, 22, 20, 14, 32, 28, 34, 8, 10, 60]): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    sm = wb.create_sheet("Summary")
    for k, v in [("Total", len(rows)), ("Sent", counts.get("sent", 0)), ("Bounced", counts.get("bounced", 0)), ("Send errors", counts.get("error", 0)), ("Not sent", counts.get("not sent", 0)), ("Run started (PT)", st.get("started", "")), ("Report type", "final with bounces" if final else "end of send")]:
        sm.append([k, v])
    name = "press_blast_report_final.xlsx" if final else "press_blast_report.xlsx"
    path = os.path.join(HERE, name); wb.save(path)
    return path, counts

def email_report(path, counts, final=False):
    from email.message import EmailMessage
    from email.utils import formatdate, make_msgid
    m = EmailMessage(); m["From"] = "Alef Press Team <press@alef.aero>"; m["To"] = ME
    m["Subject"] = ("Press blast FINAL report (with bounces)" if final else "Press blast report: sending finished") + f" - sent {counts.get('sent',0)}, bounced {counts.get('bounced',0)}, errors {counts.get('error',0)}"
    m["Date"] = formatdate(localtime=True); m["Message-ID"] = make_msgid(domain="alef.aero")
    m.set_content(f"Press blast results.\n\nSent: {counts.get('sent',0)}\nBounced: {counts.get('bounced',0)}\nSend errors: {counts.get('error',0)}\nNot sent: {counts.get('not sent',0)}\n\nDetails per recipient are in the attached spreadsheet.")
    m.add_attachment(open(path, "rb").read(), maintype="application", subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=os.path.basename(path))
    return G.send(m)

def cmd_report(args):
    path, counts = build_report(final=False); print(email_report(path, counts)); print(path, counts)

def cmd_bounces(args):
    st = load_state(); rows = json.load(open(LIST)); emails = {r["email"] for r in rows}
    started = datetime.fromisoformat(st["started"]) if st.get("started") else datetime.now(PT)
    q = f"from:mailer-daemon after:{started.strftime('%Y/%m/%d')}"
    res = G.api("GET", "messages", params={"q": q, "maxResults": 500})
    found = 0
    for m in res.get("messages", []):
        full = G.api("GET", f"messages/{m['id']}", params={"format": "full"})
        if int(full["internalDate"]) / 1000 < started.timestamp(): continue
        blob = json.dumps(full)
        for e in emails:
            if e in blob.lower():
                st["rows"].setdefault(e, {"status": "sent"}); st["rows"][e]["status"] = "bounced"; st["rows"][e]["bounce"] = full.get("snippet", "")[:200]; found += 1
    save_state(st); print("bounces found:", found)
    path, counts = build_report(final=True); print(email_report(path, counts, final=True)); print(path, counts)

if __name__ == "__main__":
    {"send": cmd_send, "report": cmd_report, "bounces": cmd_bounces}[sys.argv[1]](sys.argv[2:])
